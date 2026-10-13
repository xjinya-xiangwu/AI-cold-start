from pathlib import Path
from docx import Document
root=Path(__import__('sys').argv[1])
path=root/'docs/manual/使用手册.docx'
doc=Document(path)
# Preserve existing screenshots/tables; update text-only paragraphs without clearing drawings.
for paragraph in doc.paragraphs:
    if paragraph._p.xpath('.//w:drawing'): continue
    text=paragraph.text
    if not text: continue
    text=text.replace('v0.13.1+','v0.14.2+').replace('v0.14.0+','v0.14.2+').replace('v0.14.1+','v0.14.2+').replace('每周更新','每月内容更新').replace('3 个月周更','90天月度内容更新').replace('只新增不覆盖','精确合并，冲突保留')
    if '第 1 步' in text: text='第 1 步：下载订单专属 ZIP，完整解压到本地目录。'
    elif '第 2 步' in text: text='第 2 步：Windows 双击 install-windows.bat；macOS 双击 install-macos.command；Linux 运行 bash install-linux.sh。'
    elif text.startswith('irm '): text='Windows 命令行：powershell -NoProfile -ExecutionPolicy Bypass -File ./install.ps1 -Cli'
    elif '专属凭证' in text or '订单的专属凭证' in text: text='ZIP 的下载链接每3天刷新，过期索取卖家最新链接；已安装内容不因链接过期失效。专属链接不要公开。'
    elif '命令会弹出' in text: text='Windows 双击入口打开本地安装配置页；CLI入口直接安装。Mac/Linux当前为终端安装。'
    elif 'update-token.bat' in text: text='每月新内容发布后，下载卖家群通知/私信提供的新ZIP，解压并运行相同入口。'
    elif '重跑购买时' in text: text='更新后安装到固定目录，保留状态、冲突记录及备份。'
    elif 'enable-weekly' in text or '全自动' in text: text='默认不创建本机自动任务。卖家每72小时刷新链接；技能内容每月更新一次。'
    elif '转卖' in text: text='转载/售后范围以购买前明示条款为准；本包不能技术撤回已经下载的文件。'
    elif 'curl -fsSL' in text: text='macOS/Linux：bash ./install.sh；需要 bash、python3。原生真机兼容仍需验收。'
    if text!=paragraph.text: paragraph.text=text
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if not paragraph._p.xpath('.//w:drawing'):
                    paragraph.text=paragraph.text.replace('3 个月周更','90天月度内容更新').replace('每周更新','每月更新')
doc.save(path)
print('DOCX saved; screenshots preserved')
